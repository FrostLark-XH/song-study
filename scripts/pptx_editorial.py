"""Opt-in lyric editorial renderer. Persistent IDs own pages, native text stays editable.

Consumes visual_profile.renderVersion=3/4 and visual_assets.pages. Legacy decks keep
using L1-L8 unchanged. No lyric text is copied into the visual plan.
"""
from collections import Counter
from io import BytesIO
from pathlib import Path
import hashlib
import json
import math
import re
import unicodedata

from PIL import ImageFont
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

from scripts.atomic import atomic_replace
from scripts.pptx_fonts import resolve_all, resolve_font_file
from scripts.pptx_furigana import align_reading, reset_reading_warnings, get_reading_warnings
from scripts.pptx_layouts import _add_run
from scripts.pptx_motifs import normalize_spec, composite_motifs
from scripts.pptx_palette import contrast_ratio, rgb

W, H = 959.976, 540.0
MARGIN, TOP, BOTTOM = 64.0, 92.0, 474.0
LAYOUTS = {"left", "center", "right", "stagger"}
_FONT_CACHE = {}


def visual_fingerprint(data):
    payload={'profile':data.get('visual_profile',{}),'assets':data.get('visual_assets',{})}
    return hashlib.sha256(json.dumps(payload,ensure_ascii=False,sort_keys=True,separators=(',',':')).encode()).hexdigest()[:16]


def _hex(value):
    return isinstance(value, str) and bool(re.fullmatch(r"[0-9A-Fa-f]{6}", value))


def validate_plan(d, meta):
    pages = (d.get("visual_assets") or {}).get("pages")
    if not isinstance(pages, list) or not pages:
        raise ValueError("renderVersion=3 requires nonempty visual_assets.pages")
    records = meta["lines"]
    expected = [r["id"] for r in records]
    if len(set(expected)) != len(expected):
        raise ValueError("Duplicate source line IDs")
    by_id = {r["id"]: r for r in records}
    used, page_ids = [], []
    for p in pages:
        if not isinstance(p, dict) or not p.get("id"):
            raise ValueError("Every page needs a persistent page id")
        page_ids.append(p["id"])
        ids = p.get("line_ids", [])
        if not isinstance(ids, list) or not 1 <= len(ids) <= 3:
            raise ValueError(f"{p['id']}: support 1–3 source lines per page; split at a semantic boundary")
        if any(x not in by_id for x in ids):
            raise ValueError(f"{p['id']}: unknown line ID")
        if len({by_id[x]["section"] for x in ids}) != 1:
            raise ValueError(f"{p['id']}: a page cannot cross sections")
        if p.get("layout", "left") not in LAYOUTS:
            raise ValueError(f"{p['id']}: unknown layout")
        if p.get("typeface", "serif") not in ("serif", "sans"):
            raise ValueError(f"{p['id']}: typeface must be serif or sans")
        for field, default, lower, upper in (("font_size", 47, 32, 142), ("vertical", .45, 0, 1), ("gap", 36, 0, 120)):
            value = p.get(field, default)
            if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value) or not lower <= value <= upper:
                raise ValueError(f"{p['id']}: {field} must be between {lower} and {upper}")
        for field in ("background", "ink", "secondary", "accent"):
            if not _hex(p.get(field)):
                raise ValueError(f"{p['id']}: {field} must be six hex digits")
        if contrast_ratio(p["ink"], p["background"]) < 7:
            raise ValueError(f"{p['id']}: primary text contrast below 7:1")
        if contrast_ratio(p["secondary"], p["background"]) < 4.5:
            raise ValueError(f"{p['id']}: annotation contrast below 4.5:1")
        if not p.get("reason"):
            raise ValueError(f"{p['id']}: design reason required")
        for spec in p.get("motifs", []):
            if normalize_spec(spec, 1) is None or not spec.get("reason"):
                raise ValueError(f"{p['id']}: invalid motif or missing reason")
        used.extend(ids)
    if len(set(page_ids)) != len(page_ids):
        raise ValueError("Duplicate page IDs")
    if used != expected:
        raise ValueError("Page plan must cover each lyric occurrence once, in source order")
    if (d.get('visual_profile') or {}).get('renderVersion') == 4:
        from scripts.pptx_scene import validate_scene_plan
        validate_scene_plan(d)
    return pages, by_id


def _font(role, size):
    key = role, size
    if key not in _FONT_CACHE:
        _FONT_CACHE[key] = ImageFont.truetype(resolve_font_file(role), round(size * 4))
    return _FONT_CACHE[key]


def width(text, role, size):
    return _font(role, size).getlength(text) / 4


def textbox(slide, text, x, y, w, h, family, size, color, align="left", bold=False, name=""):
    box = slide.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
    box.name = name or "text"
    tf = box.text_frame
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    p = tf.paragraphs[0]
    p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[align]
    p.space_before = p.space_after = Pt(0)
    p.line_spacing = 1.0
    run = _add_run(p, text, family, family, Pt(size), color)
    run.font.bold = bold
    return box


def wrap_text(text, role, size, max_width):
    if not text:
        return []
    tokens = text.split(" ") if role == "romaji" else list(text)
    sep = " " if role == "romaji" else ""
    result, line = [], ""
    for token in tokens:
        candidate = line + sep + token if line else token
        if width(candidate, role, size) > max_width and line:
            result.append(line)
            line = token
        else:
            line = candidate
    if line:
        result.append(line)
    if any(width(line, role, size) > max_width for line in result):
        raise ValueError("Unbreakable annotation exceeds available width")
    return result


def wrap_units(units, size, max_width, jp_role="jp", avoid_short_tail=False):
    """Only wrap between tokenizer units; never duplicate a ruby across splits."""
    result, row, used = [], [], 0.0
    for unit in units:
        uw = width(unit["base"], jp_role, size)
        if uw > max_width:
            raise ValueError("A ruby unit cannot fit; split into another page or lower size")
        if row and used + uw > max_width:
            result.append(row)
            row, used = [], 0.0
        row.append(unit)
        used += uw
    if row:
        result.append(row)
    # Narrow illustrated columns can leave a single kana alone. Move whole
    # annotated units, preserving source order and ruby ownership. Opt-in for
    # v4 so existing v3 pagination does not silently change.
    if avoid_short_tail and len(result)>1:
        previous,tail=result[-2],result[-1]
        if len(''.join(u['base'] for u in tail).strip())==1:
            while len(previous)>1 and len(''.join(u['base'] for u in tail).strip())<3:
                if sum(width(u['base'],jp_role,size) for u in [previous[-1],*tail])>max_width:break
                tail.insert(0,previous.pop())
    return result


def prepare_line(record, size, max_width, jp_role="jp", avoid_short_tail=False):
    units = align_reading(record["jp"], record.get("romaji", ""))
    jp_rows = wrap_units(units, size, max_width, jp_role, avoid_short_tail)
    ro_rows = wrap_text(record.get("romaji", ""), "romaji", 16, max_width)
    cn_rows = wrap_text(record.get("zh", ""), "cn", 18, max_width)
    height = len(jp_rows) * (size * 1.24 + 18) + 8 + len(ro_rows) * 22 + len(cn_rows) * 25
    return {"record": record, "jp": jp_rows, "ro": ro_rows, "cn": cn_rows, "height": height}


def draw_line(slide, block, size, x, y, max_width, align, colors, fonts, jp_role="jp"):
    record = block["record"]
    for row_index, units in enumerate(block["jp"]):
        total_width = sum(width(u["base"], jp_role, size) for u in units)
        origin = x + (max_width-total_width)/2 if align == "center" else x+max_width-total_width if align == "right" else x
        text = "".join(u["base"] for u in units)
        # A single base textbox preserves CJK kerning and an editable logical line.
        # Keep a small text-engine tolerance even at right alignment. Clipping
        # this allowance at the content margin can wrap the final CJK glyph in
        # LibreOffice despite the slide's no-wrap setting.
        textbox(slide, text, origin, y+18, min(W-20-origin, total_width+size*.5), size*1.24, fonts[jp_role], size,
                colors["ink"], name=f"lyric:{record['id']}:{row_index}")
        cursor = origin
        for u in units:
            uw = width(u["base"], jp_role, size)
            ruby = u.get("ruby", "")
            if ruby:
                rs = 12.0
                while rs > 8 and width(ruby, "furigana", rs) > uw:
                    rs -= 0.5
                rw = width(ruby, "furigana", rs)
                if rw > uw + 0.5:
                    raise ValueError(f"Ruby does not fit for {record['id']}")
                textbox(slide, ruby, cursor, y, uw, 16, fonts["furigana"], rs,
                        colors["secondary"], "center", name=f"ruby:{record['id']}")
            cursor += uw
        y += size*1.24+18
    y += 8
    for text in block["ro"]:
        textbox(slide, text, x, y, max_width, 22, fonts["romaji"], 16,
                colors["secondary"], align, name=f"romaji:{record['id']}")
        y += 22
    for text in block["cn"]:
        textbox(slide, text, x, y, max_width, 25, fonts["cn"], 18,
                colors["secondary"], align, name=f"translation:{record['id']}")
        y += 25
    return y


def background(slide, page, seed):
    from PIL import Image
    specs = [normalize_spec(s, seed) for s in page.get("motifs", [])]
    if specs:
        # Existing procedural engine; the left text zone stays clear in the sample.
        canvas = Image.new("RGB", (1920, 1080), "#"+page["background"])
        import numpy as np
        result = Image.fromarray(composite_motifs(np.asarray(canvas), specs).astype("uint8"))
        buf = BytesIO()
        result.save(buf, format="PNG")
        buf.seek(0)
        slide.shapes.add_picture(buf, 0, 0, Pt(W), Pt(H)).name = "background"
    else:
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = rgb(page["background"])


def _names(d):
    rows = dict(d.get("info_rows", []))
    cfg = d.get("visual_profile", {}).get("cover", {})
    artist = cfg.get("artist") or re.split(r"[（(]", rows.get("演唱者", ""), 1)[0].strip()
    title = cfg.get("title") or d.get("title", "song").replace("歌曲学习：", "")
    if artist and " — " in title and title.rsplit(" — ", 1)[1].strip() == artist:
        title = title.rsplit(" — ", 1)[0]
    return title, artist


def render_cover(prs, d, fonts):
    cfg = d["visual_profile"].get("cover", {})
    page = {"background":"163AA6", "ink":"FFFFFF", "secondary":"DDE6FF", "motifs":[]}
    page.update(cfg)
    for key in ("background", "ink", "secondary"):
        if not _hex(page[key]):raise ValueError("Invalid cover color")
    title, artist = _names(d)
    if cfg.get("typeface", "serif") not in ("serif", "sans"):
        raise ValueError("Cover typeface must be serif or sans")
    jp_role = "gothic" if cfg.get("typeface") == "sans" else "jp"
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    is_scene = d['visual_profile'].get('renderVersion') == 4
    if is_scene:
        from scripts.pptx_scene import add_scene
        add_scene(slide,d['visual_assets'],page,W,H)
        zx,zy,zw,zh=cfg['text_zone']
        x,y,available=zx*W,zy*H,zw*W
    else:
        background(slide, page, 1)
        x,y,available=MARGIN,123,W-2*MARGIN
    textbox(slide, "LYRIC STUDY", MARGIN, 38, 300, 20, fonts["romaji"], 11, page["secondary"])
    size = 142
    while width(title, jp_role, size) > available and size > 30:size-=2
    textbox(slide, title, x, y+42, available, 170, fonts[jp_role], size, page["ink"], name="cover:title")
    textbox(slide, artist, x+4, y+226, available-4 if is_scene else 650, 36, fonts["romaji"], 25, page["secondary"])
    subtitle = cfg.get("subtitle", "")
    if subtitle:textbox(slide, subtitle, x+4, y+282, available-4, 28, fonts["cn"], 17, page["secondary"])


def build_editorial(d, meta, out_path, preview=False):
    if str(d.get("language", "")).strip().lower() != "ja":
        raise ValueError("Editorial v3 currently supports ja only; use legacy renderer for other languages")
    pages, records = validate_plan(d, meta)
    from scripts.visual_plan import check_director
    is_scene=d['visual_profile'].get('renderVersion')==4
    planning_warnings = check_director(d, meta,required=is_scene)
    for warning in planning_warnings:
        print(f"  [plan] {warning}")
    fonts = resolve_all()
    reset_reading_warnings()
    prs = Presentation()
    prs.slide_width, prs.slide_height = Pt(W), Pt(H)
    render_cover(prs, d, fonts)
    selected = pages
    if preview:
        # Preview contains actual pages in narrative order, never invented layouts.
        chosen = set(d["visual_profile"].get("previewPageIds", []))
        selected = [p for p in pages if p["id"] in chosen] if chosen else pages[:6]
        if not selected:raise ValueError("previewPageIds did not match any page")
    manifest = []
    title, artist = _names(d)
    for page in selected:
        source_index = pages.index(page)+1
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        seed = int(hashlib.sha256(page["id"].encode()).hexdigest()[:8], 16)
        if is_scene:
            from scripts.pptx_scene import add_scene
            add_scene(slide,d['visual_assets'],page,W,H)
        else:
            background(slide, page, seed)
        lines = [records[i] for i in page["line_ids"]]
        layout = page.get("layout", "left")
        if is_scene:
            zx,zy,zw,zh=page['text_zone']
            text_x,text_top,text_bottom=zx*W,zy*H,(zy+zh)*H
            available=zw*W-(50 if layout=='stagger' else 0)
        else:
            text_x,text_top,text_bottom=MARGIN,TOP,BOTTOM
            available = W-2*MARGIN-(50 if layout=="stagger" else 0)
        jp_role = "gothic" if page.get("typeface") == "sans" else "jp"
        size = float(page.get("font_size", {1:62, 2:47, 3:34}[len(lines)]))
        gap = float(page.get("gap", 12 if len(lines)==3 else 36))
        while True:
            blocks = [prepare_line(r, size, available, jp_role, avoid_short_tail=is_scene) for r in lines]
            height = sum(b["height"] for b in blocks)+gap*(len(lines)-1)
            if height <= text_bottom-text_top:break
            size -= 1
            if size < 32:raise ValueError(f"{page['id']}: vertical overflow; use fewer lyric lines")
        y = text_top + (text_bottom-text_top-height)*float(page.get("vertical", .45))
        section = lines[0]["section"].strip("[]")
        textbox(slide, section.upper(), MARGIN, 35, 630, 20, fonts["romaji"], 11,
                page["secondary"], name="section")
        textbox(slide, f"{source_index:02d} / {len(pages):02d}", W-MARGIN-110, 35, 110, 20,
                fonts["romaji"], 11, page["secondary"], "right", name="page")
        for idx, block in enumerate(blocks):
            x = text_x + (50 if layout=="stagger" and idx else 0)
            align = layout if layout in ("center", "right") else "left"
            y = draw_line(slide, block, size, x, y, available, align, page, fonts, jp_role)+gap
        textbox(slide, f"{title} / {artist}", MARGIN, 504, W-2*MARGIN, 17,
                fonts["cn"], 10, page["secondary"], name="footer")
        slide.notes_slide.notes_text_frame.text = json.dumps({
            "page_id":page["id"], "line_ids":page["line_ids"],
            "role":page.get("role"), "reason":page["reason"],
            "phase_id":page.get("phase_id"), "transition":page.get("transition"),
            "repeat_of":page.get("repeat_of"), "repeat_reason":page.get("repeat_reason"),
            "visual_event":page.get('visual_event'),"lyric_basis":page.get('lyric_basis'),
            "visual_link":page.get('visual_link'),"scene":page.get('scene'),
            "sources":d.get("verification",{}).get("lyrics",{}).get("sources",[])
        }, ensure_ascii=False)
        manifest.append({"slide":len(prs.slides),"page_id":page["id"],
                         "line_ids":page["line_ids"],"font_size":size,"layout":layout,
                         "background":page["background"],"reason":page["reason"],
                         "typeface":page.get("typeface","serif"),"role":page.get("role"),
                         "phase_id":page.get("phase_id"),"transition":page.get("transition")})
        if is_scene:
            from scripts.pptx_scene import snapshot
            manifest[-1].update({'scene':page['scene'],'text_zone':page['text_zone'],
                'visual_link':page['visual_link'],'visual_event':page['visual_event'],
                'rendered_state':snapshot(d['visual_assets'],page['scene'])})
    # Bounds check excludes only the full-canvas background.
    for i, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            if shape.name == "background":continue
            if shape.left<0 or shape.top<0 or shape.left+shape.width>Pt(W)+Pt(1) or shape.top+shape.height>Pt(H)+Pt(1):
                raise ValueError(f"Slide {i}: out-of-bounds {shape.name}")
    atomic_replace(out_path, lambda tmp: prs.save(tmp))
    report = {"renderer":"scene-v4" if is_scene else "editorial-v3", "data_fingerprint":meta["fingerprint"],
              "pptx_sha256":hashlib.sha256(Path(out_path).read_bytes()).hexdigest(),
              "visual_fingerprint":visual_fingerprint(d),
              "preview":preview, "source_line_count":len(records), "slides":len(prs.slides),
              "fonts":fonts, "pages":manifest, "reading_warnings":get_reading_warnings(),
              "planning_warnings":planning_warnings,
              "render_status":"not_run", "visual_status":"not_reviewed"}
    Path(out_path).with_suffix(".manifest.json").write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding="utf-8")
    Path(out_path).with_suffix(".reading.txt").write_text("\n".join(get_reading_warnings()) or "No alignment warnings; linguistic verification is separate.",encoding="utf-8")
    print(f"OK: {out_path}; {len(prs.slides)} slides; {sum(len(p['line_ids']) for p in selected)} source lines")
    return report
